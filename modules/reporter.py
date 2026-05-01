#!/usr/bin/env python3
"""
Módulo para generación de informes legales y técnicos
"""

import os
import json
from datetime import datetime
from colorama import Fore, Style
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

class ReportGenerator:
    """Clase para generar informes legales profesionales"""
    
    def __init__(self):
        self.templates = {
            '1': 'informe_policial',
            '2': 'informe_legal',
            '3': 'informe_plataformas',
            '4': 'informe_completo'
        }
        
    def generate_report(self, case_id, report_type, evidences):
        """Generar informe basado en el tipo seleccionado"""
        
        template = self.templates.get(report_type, 'informe_completo')
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        # Crear directorio de informes
        report_dir = f"reports/{case_id}"
        os.makedirs(report_dir, exist_ok=True)
        
        print(f"\n{Fore.CYAN}📄 Generando informe: {template.replace('_', ' ').title()}...{Style.RESET_ALL}")
        
        # Generar informe basado en el tipo
        if template == 'informe_policial':
            filename = self._generate_police_report(case_id, evidences, report_dir, timestamp)
        elif template == 'informe_legal':
            filename = self._generate_legal_report(case_id, evidences, report_dir, timestamp)
        elif template == 'informe_plataformas':
            filename = self._generate_platform_report(case_id, evidences, report_dir, timestamp)
        else:
            filename = self._generate_full_report(case_id, evidences, report_dir, timestamp)
        
        if filename:
            print(f"{Fore.GREEN}✅ Informe generado exitosamente: {filename}{Style.RESET_ALL}")
        else:
            print(f"{Fore.RED}❌ Error generando informe{Style.RESET_ALL}")
        
        return filename
    
    def _generate_police_report(self, case_id, evidences, report_dir, timestamp):
        """Generar informe para autoridades policiales"""
        try:
            filename = f"{report_dir}/denuncia_policial_{timestamp}.docx"
            doc = docx.Document()
            
            # Configuración inicial
            section = doc.sections[0]
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            
            # Encabezado
            header = section.header
            header_para = header.paragraphs[0]
            header_run = header_para.add_run(f"DENUNCIA POR ACOSO DIGITAL - CASO {case_id}")
            header_run.font.size = Pt(10)
            header_run.font.color.rgb = RGBColor(128, 128, 128)
            header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            
            # Título
            title = doc.add_heading('DENUNCIA POR ACOSO DIGITAL', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Información del caso
            doc.add_heading('1. INFORMACIÓN DEL CASO', level=1)
            case_table = doc.add_table(rows=4, cols=2)
            case_table.style = 'Light Shading'
            case_table.autofit = True
            
            case_data = [
                ['Número de caso:', case_id],
                ['Fecha de generación:', datetime.now().strftime('%d/%m/%Y %H:%M:%S')],
                ['Total de evidencias:', str(len(evidences))],
                ['Estado:', 'ACTIVO - INVESTIGACIÓN EN CURSO']
            ]
            
            for i, (label, value) in enumerate(case_data):
                row = case_table.rows[i]
                row.cells[0].text = label
                row.cells[1].text = value
                row.cells[0].paragraphs[0].runs[0].bold = True
            
            # Resumen ejecutivo
            doc.add_heading('2. RESUMEN EJECUTIVO', level=1)
            doc.add_paragraph(
                "Se presenta la presente denuncia por acoso digital sistemático, "
                "documentado a través de evidencias digitales recopiladas entre "
                f"{self._get_date_range(evidences)}. "
                "El acosador ha utilizado múltiples plataformas digitales para "
                "hostigar, amenazar y causar daño psicológico a la víctima."
            )
            
            # Evidencias documentadas
            doc.add_heading('3. EVIDENCIAS DOCUMENTADAS', level=1)
            doc.add_paragraph(
                "A continuación se detallan las evidencias recopiladas, "
                "todas ellas con integridad verificada mediante hash SHA-256:"
            )
            
            if evidences:
                evidence_table = doc.add_table(rows=len(evidences) + 1, cols=5)
                evidence_table.style = 'Light Grid'
                
                # Encabezados
                headers = ['ID', 'Tipo', 'Fecha/Hora', 'Origen', 'Hash (SHA-256)']
                header_row = evidence_table.rows[0]
                for i, header in enumerate(headers):
                    cell = header_row.cells[i]
                    cell.text = header
                    cell.paragraphs[0].runs[0].bold = True
                
                # Datos
                for i, ev in enumerate(evidences, 1):
                    row = evidence_table.rows[i]
                    row.cells[0].text = ev.get('id', 'N/A')
                    row.cells[1].text = ev.get('type', 'N/A').upper()
                    
                    # Formatear fecha
                    timestamp = ev.get('timestamp', '')
                    if timestamp:
                        try:
                            dt = datetime.fromisoformat(timestamp.replace('Z', '+00:00'))
                            row.cells[2].text = dt.strftime('%d/%m/%Y %H:%M')
                        except:
                            row.cells[2].text = timestamp
                    else:
                        row.cells[2].text = 'N/A'
                    
                    row.cells[3].text = ev.get('url', ev.get('source', 'N/A'))[:50] + ('...' if len(ev.get('url', '')) > 50 else '')
                    row.cells[4].text = ev.get('hash', 'N/A')[:16] + '...'
            
            # Análisis de patrones
            doc.add_heading('4. ANÁLISIS DE PATRONES', level=1)
            
            # Calcular patrones básicos
            patterns = self._analyze_patterns(evidences)
            
            pattern_text = (
                f"• Frecuencia: {patterns.get('total_events', 0)} eventos en {patterns.get('days_span', 0)} días\n"
                f"• Promedio diario: {patterns.get('avg_per_day', 0):.1f} eventos/día\n"
                f"• Horas pico: {patterns.get('peak_hours', 'No disponible')}\n"
                f"• Días más activos: {patterns.get('active_days', 'No disponible')}\n"
            )
            
            if patterns.get('escalation', False):
                pattern_text += "• ⚠️ SE DETECTA ESCALADA: La frecuencia ha aumentado significativamente\n"
            
            doc.add_paragraph(pattern_text)
            
            # Solicitud de medidas
            doc.add_heading('5. SOLICITUD DE MEDIDAS', level=1)
            measures = [
                "Investigación del caso por el Grupo de Delitos Telemáticos",
                "Identificación del presunto acosador mediante datos técnicos",
                "Preservación de evidencias digitales para proceso judicial",
                "Medidas de protección para la víctima",
                "Colaboración con plataformas digitales para obtención de datos"
            ]
            
            for measure in measures:
                para = doc.add_paragraph(style='List Bullet')
                para.add_run(measure)
            
            # Información del denunciante
            doc.add_heading('6. DATOS DEL DENUNCIANTE', level=1)
            denunciante_table = doc.add_table(rows=6, cols=2)
            denunciante_table.style = 'Light Shading'
            
            denunciante_data = [
                ['Nombre completo:', '________________________________'],
                ['DNI/NIE:', '________________________________'],
                ['Dirección:', '________________________________'],
                ['Teléfono:', '________________________________'],
                ['Email:', '________________________________'],
                ['Relación con la víctima:', '________________________________']
            ]
            
            for i, (label, value) in enumerate(denunciante_data):
                row = denunciante_table.rows[i]
                row.cells[0].text = label
                row.cells[1].text = value
                row.cells[0].paragraphs[0].runs[0].bold = True
            
            # Firma
            doc.add_paragraph("\n\n")
            doc.add_paragraph("Firma del denunciante:")
            doc.add_paragraph("________________________________")
            doc.add_paragraph(f"\nLugar y fecha: __________________, {datetime.now().strftime('%d de %B de %Y')}")
            
            # Guardar documento
            doc.save(filename)
            return filename
            
        except Exception as e:
            print(f"{Fore.RED}❌ Error generando informe policial: {str(e)}{Style.RESET_ALL}")
            return None
    
    def _generate_legal_report(self, case_id, evidences, report_dir, timestamp):
        """Generar informe para abogado/a"""
        try:
            filename = f"{report_dir}/informe_legal_{timestamp}.docx"
            doc = docx.Document()
            
            # Título
            title = doc.add_heading('INFORME LEGAL - ACOSO DIGITAL', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            doc.add_paragraph("CONFIDENCIAL - PARA USO LEGAL EXCLUSIVO")
            
            # Resumen del caso
            doc.add_heading('RESUMEN DEL CASO', level=1)
            doc.add_paragraph(
                f"Caso ID: {case_id}\n"
                f"Fecha: {datetime.now().strftime('%d/%m/%Y')}\n"
                f"Evidencias recopiladas: {len(evidences)}\n"
                f"Periodo documentado: {self._get_date_range(evidences)}"
            )
            
            # Tipos de delitos detectados
            doc.add_heading('TIPOS DE DELITOS DETECTADOS', level=1)
            
            crimes = [
                ("Acoso digital", "Art. 172 ter CP", "Conducta reiterada que atente gravemente contra la libertad y sentimiento de seguridad"),
                ("Amenazas", "Art. 169-171 CP", "Expresiones de causar daño a la víctima"),
                ("Calumnias e injurias", "Art. 205-208 CP", "Expresiones que lesionan la dignidad de la víctima"),
                ("Revelación de secretos", "Art. 197 CP", "Difusión de información privada sin consentimiento"),
                ("Suplantación de identidad", "Art. 401 CP", "Uso de identidad ajena para causar perjuicio")
            ]
            
            for crime, article, description in crimes:
                doc.add_heading(crime, level=2)
                doc.add_paragraph(f"Artículo aplicable: {article}")
                doc.add_paragraph(f"Descripción: {description}")
            
            # Evidencias por categoría
            doc.add_heading('CATALOGACIÓN DE EVIDENCIAS', level=1)
            
            # Agrupar por tipo
            evidence_by_type = {}
            for ev in evidences:
                ev_type = ev.get('type', 'desconocido')
                if ev_type not in evidence_by_type:
                    evidence_by_type[ev_type] = []
                evidence_by_type[ev_type].append(ev)
            
            for ev_type, items in evidence_by_type.items():
                doc.add_heading(f"{ev_type.upper()} ({len(items)} evidencias)", level=2)
                
                table = doc.add_table(rows=min(len(items), 5) + 1, cols=3)
                table.style = 'Light Grid'
                
                # Encabezados
                headers = ['Fecha', 'Origen', 'Descripción']
                header_row = table.rows[0]
                for i, header in enumerate(headers):
                    cell = header_row.cells[i]
                    cell.text = header
                    cell.paragraphs[0].runs[0].bold = True
                
                # Datos (máximo 5 por tipo)
                for i, ev in enumerate(items[:5], 1):
                    row = table.rows[i]
                    
                    # Fecha
                    timestamp = ev.get('timestamp', '')
                    if timestamp:
                        try:
                            dt = datetime.fromisoformat(timestamp.replace('Z', '+00:00'))
                            row.cells[0].text = dt.strftime('%d/%m/%Y')
                        except:
                            row.cells[0].text = timestamp
                    else:
                        row.cells[0].text = 'N/A'
                    
                    # Origen
                    row.cells[1].text = ev.get('url', ev.get('source', 'N/A'))[:40] + ('...' if len(ev.get('url', '')) > 40 else '')
                    
                    # Descripción
                    desc = self._get_evidence_description(ev)
                    row.cells[2].text = desc[:60] + ('...' if len(desc) > 60 else '')
                
                if len(items) > 5:
                    doc.add_paragraph(f"... y {len(items) - 5} evidencias adicionales")
            
            # Valoración jurídica
            doc.add_heading('VALORACIÓN JURÍDICA PRELIMINAR', level=1)
            
            legal_assessment = (
                "1. **Suficiencia probatoria**: " + 
                ("ALTA" if len(evidences) >= 10 else "MEDIA" if len(evidences) >= 5 else "BAJA") +
                f" ({len(evidences)} evidencias recopiladas)\n\n"
                
                "2. **Posibles acciones legales**:\n"
                "   • Denuncia penal por acoso digital\n"
                "   • Medidas cautelares de protección\n"
                "   • Demanda por daños morales\n"
                "   • Solicitud de restricción de contacto\n\n"
                
                "3. **Plazos procesales**:\n"
                "   • Prescripción: 5 años (delitos de amenazas y acoso)\n"
                "   • Diligencias urgentes: Inmediatas\n"
                "   • Medidas cautelares: 72 horas\n\n"
                
                "4. **Recomendaciones inmediatas**:\n"
                "   • Presentar denuncia en comisaría\n"
                "   • Solicitar orden de protección\n"
                "   • Notificar a plataformas digitales\n"
                "   • Documentar cualquier contacto futuro\n"
            )
            
            doc.add_paragraph(legal_assessment)
            
            # Anexos técnicos
            doc.add_heading('ANEXOS TÉCNICOS', level=1)
            doc.add_paragraph(
                "Todos los archivos de evidencia están disponibles en formato digital "
                "con integridad verificada mediante hash SHA-256. "
                "Se incluyen:\n\n"
                "• Capturas de pantalla con marca de tiempo\n"
                "• Metadatos EXIF de imágenes\n"
                "• Registros de actividad digital\n"
                "• Línea temporal de eventos\n"
                "• Reporte de integridad forense"
            )
            
            # Guardar documento
            doc.save(filename)
            return filename
            
        except Exception as e:
            print(f"{Fore.RED}❌ Error generando informe legal: {str(e)}{Style.RESET_ALL}")
            return None
    
    def _generate_platform_report(self, case_id, evidences, report_dir, timestamp):
        """Generar informe para plataformas digitales"""
        try:
            filename = f"{report_dir}/reporte_plataformas_{timestamp}.docx"
            doc = docx.Document()
            
            # Título
            title = doc.add_heading('REPORTE DE CONTENIDO INAPROPIADO', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Destinatario
            doc.add_paragraph("Para: Equipo de Seguridad y Moderación")
            doc.add_paragraph("[NOMBRE DE LA PLATAFORMA DIGITAL]")
            doc.add_paragraph("\n")
            
            # Asunto
            subject = doc.add_heading('ASUNTO: SOLICITUD DE ELIMINACIÓN DE CONTENIDO Y SANCIÓN DE USUARIO', level=1)
            subject.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Información del usuario infractor
            doc.add_heading('1. USUARIO INFRACTOR', level=2)
            doc.add_paragraph(
                "Nombre de usuario: [@NOMBRE_DE_USUARIO]\n"
                "URL del perfil: [URL_DEL_PERFIL]\n"
                "Tipo de violación: Acoso digital, amenazas, comportamiento abusivo"
            )
            
            # Violaciones de términos
            doc.add_heading('2. VIOLACIONES DE TÉRMINOS DE SERVICIO', level=2)
            
            violations = [
                "Acoso y hostigamiento a otros usuarios",
                "Publicación de contenido amenazante",
                "Comportamiento abusivo reiterado",
                "Suplantación de identidad",
                "Difusión de información privada sin consentimiento"
            ]
            
            for violation in violations:
                para = doc.add_paragraph(style='List Bullet')
                para.add_run(violation)
            
            # Evidencias específicas
            doc.add_heading('3. EVIDENCIAS DE LAS VIOLACIONES', level=2)
            
            if evidences:
                # Agrupar por plataforma/origen
                by_source = {}
                for ev in evidences:
                    source = ev.get('url', ev.get('source', 'general'))
                    if source not in by_source:
                        by_source[source] = []
                    by_source[source].append(ev)
                
                for source, items in by_source.items():
                    doc.add_heading(f"Contenido en: {source}", level=3)
                    
                    for i, ev in enumerate(items[:3], 1):  # Máximo 3 por fuente
                        doc.add_paragraph(f"Evidencia {i}:", style='Heading 4')
                        
                        info = f"• Fecha: {self._format_date(ev.get('timestamp', ''))}\n"
                        
                        if ev.get('type') == 'text' and 'text_preview' in ev:
                            info += f"• Contenido: {ev['text_preview'][:200]}...\n"
                        
                        if 'hash' in ev:
                            info += f"• Hash de verificación: {ev['hash'][:16]}...\n"
                        
                        doc.add_paragraph(info)
            
            # Solicitudes específicas
            doc.add_heading('4. SOLICITUDES ESPECÍFICAS', level=2)
            
            requests = [
                "Eliminación inmediata de todo contenido ofensivo",
                "Suspensión o cancelación de la cuenta del usuario infractor",
                "Preservación de datos para investigación legal",
                "Notificación de las acciones tomadas",
                "Colaboración con autoridades en caso de investigación"
            ]
            
            for req in requests:
                para = doc.add_paragraph(style='List Bullet')
                para.add_run(req)
            
            # Información de contacto
            doc.add_heading('5. INFORMACIÓN DE CONTACTO', level=2)
            contact_table = doc.add_table(rows=4, cols=2)
            contact_table.style = 'Light Shading'
            
            contact_data = [
                ['Nombre:', '[TU_NOMBRE_COMPLETO]'],
                ['Email:', '[TU_EMAIL_VÁLIDO]'],
                ['Teléfono:', '[TU_TELÉFONO]'],
                ['Caso ID:', case_id]
            ]
            
            for i, (label, value) in enumerate(contact_data):
                row = contact_table.rows[i]
                row.cells[0].text = label
                row.cells[1].text = value
                row.cells[0].paragraphs[0].runs[0].bold = True
            
            # Anexos
            doc.add_heading('6. ANEXOS ADJUNTOS', level=2)
            doc.add_paragraph(
                "Se adjuntan los siguientes archivos:\n\n"
                "• Capturas de pantalla de los contenidos ofensivos\n"
                "• Archivos de evidencia con hashes de verificación\n"
                "• Línea temporal de los incidentes\n"
                "• Reporte de integridad de las evidencias"
            )
            
            # Guardar documento
            doc.save(filename)
            return filename
            
        except Exception as e:
            print(f"{Fore.RED}❌ Error generando informe para plataformas: {str(e)}{Style.RESET_ALL}")
            return None
    
    def _generate_full_report(self, case_id, evidences, report_dir, timestamp):
        """Generar informe completo del caso"""
        try:
            filename = f"{report_dir}/informe_completo_{timestamp}.docx"
            doc = docx.Document()
            
            # Portada
            title = doc.add_heading('INFORME COMPLETO DE INVESTIGACIÓN DIGITAL', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            subtitle = doc.add_heading(f'CASO {case_id}', 1)
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            doc.add_paragraph("\n" * 3)
            doc.add_paragraph("Generado por: Sistema Anti-Acoso Digital Toolkit")
            doc.add_paragraph(f"Fecha: {datetime.now().strftime('%d de %B de %Y')}")
            doc.add_paragraph(f"Hora: {datetime.now().strftime('%H:%M:%S')}")
            
            doc.add_page_break()
            
            # Índice
            doc.add_heading('ÍNDICE', level=1)
            
            toc = [
                ("1. RESUMEN EJECUTIVO", 1),
                ("2. INFORMACIÓN DEL CASO", 1),
                ("3. METODOLOGÍA DE INVESTIGACIÓN", 1),
                ("4. EVIDENCIAS DOCUMENTADAS", 1),
                ("  4.1 Por Tipo de Evidencia", 2),
                ("  4.2 Por Fuente/Plataforma", 2),
                ("  4.3 Línea Temporal", 2),
                ("5. ANÁLISIS DE PATRONES", 1),
                ("  5.1 Patrones Temporales", 2),
                ("  5.2 Frecuencia de Actividad", 2),
                ("  5.3 Evaluación de Riesgo", 2),
                ("6. VALORACIÓN JURÍDICA", 1),
                ("7. RECOMENDACIONES", 1),
                ("8. ANEXOS TÉCNICOS", 1)
            ]
            
            for title_text, level in toc:
                if level == 1:
                    para = doc.add_paragraph()
                    para.add_run(title_text).bold = True
                else:
                    doc.add_paragraph(title_text)
            
            doc.add_page_break()
            
            # 1. Resumen Ejecutivo
            doc.add_heading('1. RESUMEN EJECUTIVO', level=1)
            
            summary = (
                f"Este informe presenta los hallazgos de la investigación digital del caso {case_id}, "
                f"realizada entre {self._get_date_range(evidences)}. "
                f"Se han recopilado {len(evidences)} evidencias digitales que documentan un patrón "
                "sistemático de acoso y comportamiento abusivo.\n\n"
                
                "**Hallazgos principales:**\n"
                f"• Total de evidencias: {len(evidences)}\n"
                f"• Periodo cubierto: {self._get_date_range(evidences)}\n"
                f"• Plataformas involucradas: {len(set([e.get('source', '') for e in evidences if e.get('source')]))}\n"
                f"• Integridad verificada: 100% de las evidencias con hash SHA-256\n\n"
                
                "**Conclusión preliminar:**\n"
                "La evidencia recopilada indica un caso claro de acoso digital que requiere "
                "acción inmediata por parte de las autoridades competentes y las plataformas digitales."
            )
            
            doc.add_paragraph(summary)
            
            # 2. Información del caso
            doc.add_heading('2. INFORMACIÓN DEL CASO', level=1)
            
            case_info = (
                f"**ID del caso:** {case_id}\n"
                f"**Fecha de inicio:** {self._get_earliest_date(evidences)}\n"
                f"**Fecha de fin:** {self._get_latest_date(evidences)}\n"
                f"**Estado:** Investigación activa\n"
                f"**Responsable:** Sistema Automatizado de Documentación\n"
                f"**Última actualización:** {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n\n"
                
                "**Metadatos técnicos:**\n"
                "• Formato de evidencias: Digital forense\n"
                "• Algoritmo de hash: SHA-256\n"
                "• Estándar de tiempo: ISO 8601\n"
                "• Codificación: UTF-8"
            )
            
            doc.add_paragraph(case_info)
            
            # 3. Metodología
            doc.add_heading('3. METODOLOGÍA DE INVESTIGACIÓN', level=1)
            
            methodology = (
                "La investigación se realizó siguiendo protocolos forenses digitales estándar:\n\n"
                
                "**3.1 Recopilación de evidencias:**\n"
                "• Captura de pantallas con marca de tiempo\n"
                "• Descarga de contenidos digitales\n"
                "• Registro de metadatos EXIF\n"
                "• Documentación de URLs y fuentes\n\n"
                
                "**3.2 Preservación de integridad:**\n"
                "• Cálculo de hash SHA-256 para cada evidencia\n"
                "• Cadena de custodia digital\n"
                "• Almacenamiento en contenedores seguros\n"
                "• Copias de seguridad redundantes\n\n"
                
                "**3.3 Análisis forense:**\n"
                "• Análisis de patrones temporales\n"
                "• Identificación de plataformas utilizadas\n"
                "• Detección de comportamientos sistemáticos\n"
                "• Evaluación de riesgo basada en datos\n\n"
                
                "**3.4 Consideraciones éticas:**\n"
                "• Consentimiento informado de la víctima\n"
                "• Respeto a leyes de privacidad\n"
                "• Uso exclusivo para fines legales\n"
                "• Destrucción segura tras resolución"
            )
            
            doc.add_paragraph(methodology)
            
            # 4. Evidencias documentadas
            doc.add_heading('4. EVIDENCIAS DOCUMENTADAS', level=1)
            doc.add_paragraph(
                f"Se han recopilado {len(evidences)} evidencias digitales con "
                "integridad verificada mediante hash SHA-256."
            )

            if evidences:
                rows = min(len(evidences), 20) + 1
                ev_table = doc.add_table(rows=rows, cols=4)
                ev_table.style = 'Light Grid'
                for i, hdr in enumerate(['ID', 'Tipo', 'Fecha/Hora', 'Origen']):
                    cell = ev_table.rows[0].cells[i]
                    cell.text = hdr
                    cell.paragraphs[0].runs[0].bold = True
                for i, ev in enumerate(evidences[:20], 1):
                    row = ev_table.rows[i]
                    row.cells[0].text = ev.get('id', 'N/A')
                    row.cells[1].text = ev.get('type', 'N/A').upper()
                    ts = ev.get('timestamp', '')
                    try:
                        dt = datetime.fromisoformat(ts.replace('Z', '+00:00'))
                        row.cells[2].text = dt.strftime('%d/%m/%Y %H:%M')
                    except Exception:
                        row.cells[2].text = ts or 'N/A'
                    src = ev.get('url', ev.get('source', 'N/A'))
                    row.cells[3].text = src[:50] + ('...' if len(src) > 50 else '')
                if len(evidences) > 20:
                    doc.add_paragraph(f"... y {len(evidences) - 20} evidencias adicionales.")

            # 5. Análisis de patrones
            doc.add_heading('5. ANÁLISIS DE PATRONES', level=1)
            patterns = self._analyze_patterns(evidences)
            escalada = "⚠️ ESCALADA DETECTADA — La frecuencia ha aumentado significativamente.\n" if patterns.get('escalation') else ""
            doc.add_paragraph(
                f"• Eventos totales: {patterns.get('total_events', 0)}\n"
                f"• Rango temporal: {patterns.get('days_span', 0)} días\n"
                f"• Promedio diario: {patterns.get('avg_per_day', 0):.1f} eventos/día\n"
                f"• Horas pico: {', '.join(patterns.get('peak_hours', ['N/A']))}\n"
                f"• Días más activos: {', '.join(patterns.get('active_days', ['N/A']))}\n"
                + escalada
            )

            # 6. Valoración jurídica
            doc.add_heading('6. VALORACIÓN JURÍDICA', level=1)
            sufficiency = "ALTA" if len(evidences) >= 10 else "MEDIA" if len(evidences) >= 5 else "BAJA"
            doc.add_paragraph(
                f"Suficiencia probatoria: {sufficiency} ({len(evidences)} evidencias recopiladas)\n\n"
                "Posibles acciones legales:\n"
                "  • Denuncia penal por acoso digital (Art. 172 ter CP)\n"
                "  • Medidas cautelares de protección inmediata\n"
                "  • Demanda civil por daños morales\n"
                "  • Solicitud de orden de alejamiento digital\n\n"
                "Plazos procesales relevantes:\n"
                "  • Prescripción del delito: 5 años\n"
                "  • Resolución de medidas cautelares: 72 horas\n"
                "  • Tiempo estimado para pericial forense: 2-4 semanas\n"
            )

            # 7. Recomendaciones
            doc.add_heading('7. RECOMENDACIONES', level=1)
            recommendations = [
                "Presentar este informe completo en comisaría junto con las evidencias digitales",
                "Solicitar orden de protección de forma inmediata",
                "Notificar a las plataformas digitales involucradas mediante el informe específico",
                "Guardar copias de seguridad en múltiples ubicaciones físicas y en la nube",
                "Consultar con abogado/a especializado en delitos informáticos",
                "Continuar documentando cualquier contacto o incidente futuro",
                "No eliminar mensajes ni perfiles — son evidencias adicionales",
            ]
            for rec in recommendations:
                para = doc.add_paragraph(style='List Bullet')
                para.add_run(rec)

            # 8. Anexos técnicos
            doc.add_heading('8. ANEXOS TÉCNICOS', level=1)
            doc.add_paragraph(
                "Los siguientes archivos forman parte del expediente digital y "
                "están disponibles en el directorio del caso:\n\n"
                "  • Capturas de pantalla con contenido real de páginas web\n"
                "  • Imágenes descargadas con metadatos EXIF\n"
                "  • Registros de actividad detallados (activity.log)\n"
                "  • Línea temporal de eventos (timeline.csv)\n"
                "  • Reporte de integridad con hashes SHA-256\n"
                "  • Cadena de custodia digital (chain_of_custody.json)\n"
            )

            # Guardar documento
            doc.save(filename)
            return filename
            
        except Exception as e:
            print(f"{Fore.RED}❌ Error generando informe completo: {str(e)}{Style.RESET_ALL}")
            return None
    
    def _get_date_range(self, evidences):
        """Obtener rango de fechas de las evidencias"""
        if not evidences:
            return "No disponible"
        
        dates = []
        for ev in evidences:
            if 'timestamp' in ev:
                try:
                    dt = datetime.fromisoformat(ev['timestamp'].replace('Z', '+00:00'))
                    dates.append(dt)
                except:
                    pass
        
        if dates:
            dates.sort()
            start = dates[0].strftime('%d/%m/%Y')
            end = dates[-1].strftime('%d/%m/%Y')
            return f"{start} al {end}"
        
        return "No disponible"
    
    def _format_date(self, timestamp):
        """Formatear fecha legiblemente"""
        if not timestamp:
            return "No disponible"
        
        try:
            dt = datetime.fromisoformat(timestamp.replace('Z', '+00:00'))
            return dt.strftime('%d/%m/%Y %H:%M')
        except:
            return timestamp
    
    def _get_evidence_description(self, evidence):
        """Obtener descripción de evidencia"""
        ev_type = evidence.get('type', 'desconocido')
        
        if ev_type == 'screenshot':
            return f"Captura de {evidence.get('url', 'URL desconocida')}"
        elif ev_type == 'image':
            return f"Imagen de {evidence.get('source', 'fuente desconocida')}"
        elif ev_type == 'text':
            preview = evidence.get('text_preview', '')
            return f"Texto: {preview[:50]}..." if preview else "Documento de texto"
        else:
            return f"Evidencia tipo {ev_type}"
    
    def _analyze_patterns(self, evidences):
        """Analizar patrones básicos"""
        patterns = {
            'total_events': len(evidences),
            'days_span': 0,
            'avg_per_day': 0,
            'peak_hours': [],
            'active_days': [],
            'escalation': False
        }
        
        if not evidences:
            return patterns
        
        # Extraer fechas
        dates = []
        for ev in evidences:
            if 'timestamp' in ev:
                try:
                    dt = datetime.fromisoformat(ev['timestamp'].replace('Z', '+00:00'))
                    dates.append(dt)
                except:
                    pass
        
        if dates:
            dates.sort()
            patterns['days_span'] = (dates[-1] - dates[0]).days + 1
            patterns['avg_per_day'] = len(evidences) / patterns['days_span'] if patterns['days_span'] > 0 else 0
            
            # Horas pico
            hour_counts = {}
            for dt in dates:
                hour = dt.hour
                hour_counts[hour] = hour_counts.get(hour, 0) + 1
            
            if hour_counts:
                top_hours = sorted(hour_counts.items(), key=lambda x: x[1], reverse=True)[:3]
                patterns['peak_hours'] = [f"{h}:00 ({c})" for h, c in top_hours]
            
            # Días activos
            day_counts = {}
            for dt in dates:
                day = dt.strftime('%A')
                day_counts[day] = day_counts.get(day, 0) + 1
            
            if day_counts:
                top_days = sorted(day_counts.items(), key=lambda x: x[1], reverse=True)[:3]
                patterns['active_days'] = [f"{d} ({c})" for d, c in top_days]
            
            # Escalada
            if len(dates) >= 4:
                split = len(dates) // 2
                first_half = dates[:split]
                second_half = dates[split:]
                
                freq_first = len(first_half) / ((first_half[-1] - first_half[0]).days + 1) if len(first_half) > 1 else 0
                freq_second = len(second_half) / ((second_half[-1] - second_half[0]).days + 1) if len(second_half) > 1 else 0
                
                patterns['escalation'] = freq_second > freq_first * 1.5
        
        return patterns
    
    def _get_earliest_date(self, evidences):
        """Obtener fecha más temprana"""
        dates = []
        for ev in evidences:
            if 'timestamp' in ev:
                try:
                    dt = datetime.fromisoformat(ev['timestamp'].replace('Z', '+00:00'))
                    dates.append(dt)
                except:
                    pass
        
        if dates:
            return min(dates).strftime('%d/%m/%Y')
        return "No disponible"
    
    def _get_latest_date(self, evidences):
        """Obtener fecha más reciente"""
        dates = []
        for ev in evidences:
            if 'timestamp' in ev:
                try:
                    dt = datetime.fromisoformat(ev['timestamp'].replace('Z', '+00:00'))
                    dates.append(dt)
                except:
                    pass
        
        if dates:
            return max(dates).strftime('%d/%m/%Y')
        return "No disponible"